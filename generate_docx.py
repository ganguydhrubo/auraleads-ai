import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    # Normal style
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("AuraLeads AI (AuraLeads.ai)")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep blue

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Comprehensive End-to-End Product, Architecture & Deployment Handoff Specification")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Metadata Box Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Prepared For", "chetan@1xl.com / Dhrubojyoti"),
        ("Date & Status", "25 September 2026 — Live Production Release"),
        ("Live Production URL", "https://celestia-leads.vercel.app"),
        ("Source Repository", "https://github.com/ganguydhrubo/auraleads-ai")
    ]

    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FAFAFA")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for Section Headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        return h

    # Section 1
    add_heading_1("1. Executive Summary & Brand Transformation")
    p = doc.add_paragraph()
    p.add_run(
        "AuraLeads AI is an autonomous, multi-channel B2B lead generation, scraping, and outreach engine. "
        "Engineered as a high-performance enhancement over the inspected celestialeads.com platform, "
        "AuraLeads AI eliminates static mock-ups in favor of real-time AI inference, live geographic data scraping, "
        "multi-tenant workspace persistence, and high-converting marketing assets."
    )

    doc.add_paragraph(
        "Key Differentiators & Enhancements:\n"
        "• High-Converting Interactive Landing Page: Dynamic 4-stage pipeline demo, ROI value calculator, and monthly/annual pricing toggle.\n"
        "• Agentic SEO & Geo-Optimization: Structured JSON-LD schemas (SoftwareApplication, Organization, FAQPage), OpenGraph tags, and geo-targeted meta descriptors.\n"
        "• High-Speed Groq AI Inference: Integrated with Groq's ultra-fast LLM pipeline (openai/gpt-oss-20b) delivering sub-second AI hashtag validation and personalized outreach copy.\n"
        "• Multi-Tenant Relational Database: Full PostgreSQL schema running on Supabase with 8 dedicated tables and Row Level Security (RLS).\n"
        "• Live OpenStreetMap Geocoding: Interactive Leaflet boundary polygon preview connected to OpenStreetMap Nominatim.\n"
        "• Clean Production Integrations: Pre-wired pipelines ready for user-owned Meta Instagram Graph API credentials and Gmail App Passwords without fake mock data contamination."
    )

    # Section 2
    add_heading_1("2. Master Links & Production Command Center")
    links_table = doc.add_table(rows=6, cols=3)
    links_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    links_table.autofit = False

    headers = ["Service / Resource", "Location / URL", "Operational Status"]
    for j, h in enumerate(headers):
        cell = links_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    rows_data = [
        ("Live Web App (Production Alias)", "https://celestia-leads.vercel.app", "HTTP 200 OK — Active"),
        ("Latest Production Deployment", "https://celestia-leads-7oo2qt7gd-dhrubojyoti.vercel.app", "Deployed & Verified"),
        ("Vercel Project Control Panel", "https://vercel.com/dhrubojyoti/celestia-leads", "Connected with Production Secrets"),
        ("GitHub Source Code Repository", "https://github.com/ganguydhrubo/auraleads-ai", "Synced ('main' branch)"),
        ("Supabase Database & SQL Editor", "https://supabase.com/dashboard/project/tqvvqskscawuifjkaqby", "Online (PostgreSQL 17.6)")
    ]

    for i, row_data in enumerate(rows_data, start=1):
        row = links_table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if j == 0:
                cell.width = Inches(2.2)
            elif j == 1:
                cell.width = Inches(3.0)
            else:
                cell.width = Inches(1.3)
            bg = "FFFFFF" if i % 2 != 0 else "F9FAFB"
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Green

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3
    add_heading_1("3. System Architecture & Component Mapping")
    doc.add_paragraph(
        "AuraLeads AI is engineered with Next.js 14 App Router, modern TypeScript, and Tailwind CSS. "
        "The architecture is segmented into three decoupled tiers:"
    )
    doc.add_paragraph(
        "1. Presentation & Interaction Tier:\n"
        "   - Landing Page (app/page.tsx & components/landing/LandingPage.tsx): High-conversion sales vehicle with interactive widgets.\n"
        "   - Multi-Tenant Authentication (app/login/page.tsx, app/signup/page.tsx): Workspace isolation and session provisioning.\n"
        "   - Web App Shell (app/app/page.tsx): Single-page dynamic view-switching container managing 15 modular views.\n"
        "   - Interactive Shell Components: 3-tier collapsible sidebar, quota countdown pills, 5-step onboarding banner, 15-step walkthrough tour, and persistent support chat.\n\n"
        "2. Serverless API Routing Tier (app/api/*):\n"
        "   - POST /api/generate/hashtags: Real-time Groq LLM inference generating discovery tags.\n"
        "   - POST /api/generate/message: Groq LLM personalization synthesizing Instagram DMs and cold emails.\n"
        "   - GET /api/maps/search: OpenStreetMap Nominatim proxy returning GPS coordinates and boundary GeoJSON polygons.\n"
        "   - POST /api/maps/scrape: Business extraction and executive contact enrichment.\n"
        "   - GET & POST /api/state: Supabase-backed workspace state synchronizer with memory fallback.\n"
        "   - GET & POST /api/webhooks/instagram: Meta Graph API webhook endpoint supporting hub.challenge and HMAC SHA-256 signatures.\n\n"
        "3. Cloud Infrastructure & Persistence Tier:\n"
        "   - Vercel Edge Network: Worldwide CDN delivery, SSR caching, and serverless function deployment.\n"
        "   - Groq Cloud: High-speed LPUs running openai/gpt-oss-20b for rapid natural language generation.\n"
        "   - Supabase Managed PostgreSQL: High-availability database hosting multi-tenant schemas and Row Level Security."
    )

    # Section 4
    add_heading_1("4. Screen-by-Screen Product Inventory (All 15 Views)")
    views = [
        ("Hashtags Setup", "Business niche definition, geo-targeting specifications, and live Groq AI hashtag generation with validation scores."),
        ("Hashtags Leads Queue", "Discovered Instagram profiles, weekly quota countdowns, matched vs blocked filtering, and bulk CSV export."),
        ("Competitors Setup", "Addition of up to 5 Instagram competitor handles with 7-day anti-abuse cycle locking mechanisms."),
        ("Competitor Leads Queue", "Ingestion feed of competitor followers with qualification scoring and 1-click DM triggering."),
        ("Smart Lead Filters", "Granular audience filtering: content themes, follower brackets, location boundaries, language flags, and contact requirement toggles."),
        ("Google Maps Discovery", "Interactive Leaflet map previewing OpenStreetMap geographic polygon boundaries and local query parameters."),
        ("Google Maps Leads", "Revealed business entities, decision-maker contact lookup (Owner/CEO/Founder), direct telephone, and email."),
        ("Unified 3-Pane Inbox", "Omnichannel communication hub dividing Leads from Users, live DM/email switcher, and AI Auto-Reply rules modal."),
        ("Outreach Campaigns", "Automated multi-step outbound sequences, manual CSV campaign dispatcher, and browser extension automated sender."),
        ("Copywriting Templates", "Pre-configured DM pitches, cold email templates, personalization merge tags, and signature preview."),
        ("Analytics Dashboard", "Yield analytics, conversion funnel charts, reply rates, and hashtag/competitor performance metrics."),
        ("Billing & Plans", "Silver ($20/mo), Gold ($50/mo), Platinum ($100/mo) tiers, comments add-on, and integrated PayPal checkout modal."),
        ("Integrations & Settings", "BYO Meta Developer App credentials, webhook verification keys, and Gmail App Password SMTP connectors."),
        ("Admin Console", "Worker scraper account pool management, Meta application pool rotation, and waitlist capacity monitors."),
        ("Product Guidance Shell", "5-step onboarding checklist, 15-step interactive product tour, and sliding customer support drawer.")
    ]

    for title, desc in views:
        p = doc.add_paragraph()
        r1 = p.add_run(f"• {title}: ")
        r1.font.bold = True
        r2 = p.add_run(desc)

    # Section 5
    add_heading_1("5. Supabase Database Schema Specifications")
    doc.add_paragraph(
        "The Supabase project (tqvvqskscawuifjkaqby) is configured with PostgreSQL 17.6. "
        "All 8 production tables are actively deployed with foreign key cascades and Row Level Security (RLS):"
    )

    schema_table = doc.add_table(rows=9, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    schema_table.autofit = False

    s_headers = ["Table Name", "Key Columns & Data Types", "Business Purpose"]
    for j, h in enumerate(s_headers):
        cell = schema_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    tables_data = [
        ("workspaces", "id (UUID PK), name (TEXT), domain (TEXT), plan (TEXT), trial_ends_at (TIMESTAMPTZ), credits_remaining (INT)", "Multi-tenant tenant isolation and billing tracking"),
        ("users", "id (UUID PK), workspace_id (FK), email (TEXT UNIQUE), password_hash (TEXT), full_name (TEXT), role (TEXT)", "Authentication credentials and user access levels"),
        ("leads", "id (UUID PK), workspace_id (FK), platform (TEXT), username, name, email, phone, website, ai_score, metadata (JSONB)", "Prospect repository from Instagram & Google Maps"),
        ("campaigns", "id (UUID PK), workspace_id (FK), name (TEXT), status (TEXT), channel (TEXT), sent_count (INT), replied_count (INT)", "Outreach sequence configuration and delivery metrics"),
        ("messages", "id (UUID PK), workspace_id (FK), lead_id (FK), sender (TEXT), channel (TEXT), content (TEXT), is_read (BOOL)", "Omnichannel conversational history and AI messages"),
        ("hashtags", "id (UUID PK), workspace_id (FK), tag (TEXT), validation_score (INT), posts_count (TEXT), relevance_score (INT)", "AI-generated discovery hashtags and performance scores"),
        ("competitors", "id (UUID PK), workspace_id (FK), handle (TEXT), cycle_reset_at (TIMESTAMPTZ)", "Instagram competitor tracking and 7-day ingestion locks"),
        ("workspace_settings", "id (UUID PK), workspace_id (FK UNIQUE), meta_app_id, meta_app_secret, meta_access_token, gmail_account", "Encrypted BYO API keys and mail credentials")
    ]

    for i, row_data in enumerate(tables_data, start=1):
        row = schema_table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if j == 0:
                cell.width = Inches(1.5)
            elif j == 1:
                cell.width = Inches(3.2)
            else:
                cell.width = Inches(1.8)
            bg = "FFFFFF" if i % 2 != 0 else "F9FAFB"
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(9.0)
            if j == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 6
    add_heading_1("6. Production Environment Variables")
    doc.add_paragraph("The following production secrets are securely provisioned in Vercel:")
    
    env_table = doc.add_table(rows=6, cols=3)
    env_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    env_table.autofit = False

    e_headers = ["Variable Name", "Target Environment", "Description & Usage"]
    for j, h in enumerate(e_headers):
        cell = env_table.rows[0].cells[j]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    env_rows = [
        ("GROQ_API_KEY", "Production, Preview", "Groq Cloud high-speed LLM inference key (openai/gpt-oss-20b)"),
        ("NEXT_PUBLIC_SUPABASE_URL", "Production, Preview, Dev", "Supabase project REST endpoint (https://tqvvqskscawuifjkaqby.supabase.co)"),
        ("NEXT_PUBLIC_SUPABASE_ANON_KEY", "Production, Preview, Dev", "Public anon client key for browser and client queries"),
        ("SUPABASE_SERVICE_ROLE_KEY", "Production (Sensitive Secret)", "Administrative service role key for backend database transactions"),
        ("META_VERIFY_TOKEN", "Production, Preview", "Webhook verification token for Meta Instagram Graph handshake")
    ]

    for i, row_data in enumerate(env_rows, start=1):
        row = env_table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if j == 0:
                cell.width = Inches(2.5)
            elif j == 1:
                cell.width = Inches(1.8)
            else:
                cell.width = Inches(2.2)
            bg = "FFFFFF" if i % 2 != 0 else "F9FAFB"
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(9.0)
            if j == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 7
    add_heading_1("7. Security, Credential Rotation & Scaling Guidelines")
    doc.add_paragraph(
        "To ensure maximum enterprise-grade security following the completion of this build:\n"
        "1. Token Rotation: The temporary Vercel token (vcp_...) and Supabase token (sbp_...) used during automated provisioning should now be revoked from your respective account settings dashboards.\n"
        "2. Git Push Protection: The GitHub repository (github.com/ganguydhrubo/auraleads-ai) is protected against secret leaks. Ensure all developers use .env.local for local overrides.\n"
        "3. Meta App Verification: When connecting your live Instagram Business account, configure the webhook callback URL to https://celestia-leads.vercel.app/api/webhooks/instagram using verify token auraleads_wh_token_88921.\n"
        "4. Gmail Multi-Inbox: In Settings -> Integrations, instruct users to generate Google 16-character App Passwords rather than standard account passwords for secure SMTP outreach."
    )

    doc.add_paragraph(
        "Sign-Off: Platform build, database migration, API test suite, and deployment verification are 100% complete and fully operational."
    )

    output_path = "C:\\Users\\Dhrubo\\.gemini\\antigravity\\scratch\\celestia-leads\\AuraLeads_AI_Complete_End_to_End_Handoff.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == '__main__':
    create_document()
